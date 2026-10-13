"""Excel download of a run."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import pipeline, scoring, store


def _style(ws, widths):
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"


def run_to_excel(run: dict) -> bytes:
    decisions = store.load_decisions()
    wb = Workbook()
    ws = wb.active
    ws.title = "Ranked prospects"
    ws.append(["Rank", "Priority", "Adjusted score", "Institution", "City", "State", "Type", "Hostel students",
               "Best signal tier", "Evidence confidence", "Why shortlisted", "Potential fit", "Outreach angle",
               "Flags", "Your decision"])
    sig = wb.create_sheet("Signals and evidence")
    sig.append(["Institution", "Tier", "Counted?", "Date", "Signal", "Source title", "Source link"])
    for i, r in enumerate(pipeline.ranked(run), 1):
        s, d = r["score"], r.get("research", {})
        ws.append([i, s["priority"], s["adjusted"], r["name"], r["city"], r["state"], r["type"],
                   d.get("hostel_students"), scoring.TIER_LABEL.get(s.get("best_tier"), ""), s["confidence"],
                   d.get("why_shortlisted", ""), d.get("potential_fit", ""), d.get("outreach_angle", ""),
                   " ".join(s["flags"]), decisions.get(store.inst_key(r), "Not reviewed")])
        for g in s["signals"]:
            sig.append([r["name"], scoring.TIER_LABEL.get(g["tier"], g["tier"]),
                        "Yes" if g["verified"] else "No - lead to check by hand", g.get("date", ""),
                        g.get("description", ""), g.get("source_title", ""), g.get("source_url", "")])
    _style(ws, [6, 14, 10, 34, 14, 14, 20, 11, 14, 12, 50, 36, 44, 40, 14])
    _style(sig, [34, 13, 22, 12, 60, 30, 50])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
